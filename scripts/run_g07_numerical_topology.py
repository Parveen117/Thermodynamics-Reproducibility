from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import urllib.parse
import urllib.request
from collections import Counter, deque
from pathlib import Path
from typing import Any, Iterable


USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "Chrome/151.0 Safari/537.36 Thermodynamics-Reproducibility-G07/1.0"
)

CATEGORY_PATTERNS: dict[str, tuple[str, ...]] = {
    "temperature": (
        r"^t$",
        r"\bt\s*[=(]",
        r"\bt\s*\(\s*k\s*\)",
        r"\btemperature\b",
        r"\belectron(?:ic)?\s+temperature\b",
        r"\bfermi\s+temperature\b",
        r"\b\d+(?:\.\d+)?\s*k\b",
    ),
    "density": (
        r"^n$",
        r"\bn\s*[=(]",
        r"\bn_?min\b",
        r"\bcarrier\s+density\b",
        r"\bcharge\s+density\b",
        r"\bdoping(?:\s+regime)?\b",
        r"\bvg\b",
        r"cm\s*[-^]?\s*2",
    ),
    "electrical": (
        r"\belectrical\s+conductiv",
        r"\bconductance\b",
        r"\bresistance\b",
        r"\bsigma(?:_q)?\b",
        r"σ",
        r"^g(?:\s|\(|$)",
    ),
    "thermal": (
        r"\bthermal\s+conductiv",
        r"\bkappa(?:_e)?\b",
        r"κ",
        r"\blorenz\b",
        r"\bwiedemann",
        r"\bl\s*/\s*l0\b",
    ),
    "derived": (
        r"\bviscosity\b",
        r"\bentropy\b",
        r"\benthalpy\b",
        r"\bmean\s+free\s+path\b",
        r"\bknudsen\b",
        r"\beta(?:_th)?\b",
    ),
    "uncertainty": (
        r"\berr(?:or)?\b",
        r"\buncertainty\b",
        r"\bstandard\s+deviation\b",
        r"\bconfidence\s+interval\b",
    ),
    "units_or_scaling": (
        r"\bk\b",
        r"cm\s*[-^]?\s*2",
        r"\bsiemens\b",
        r"\bohm\b",
        r"\bw\s*m",
        r"\bj\s*m",
        r"\be\^?2\s*/\s*h\b",
        r"\bl\s*/\s*l0\b",
        r"\b1e\d+\b",
        r"\b10\^",
    ),
}
COMPILED_PATTERNS = {
    group: tuple(re.compile(pattern, re.I) for pattern in patterns)
    for group, patterns in CATEGORY_PATTERNS.items()
}


def sanitize_url(url: str) -> str:
    parts = urllib.parse.urlsplit(url.strip())
    path = urllib.parse.quote(parts.path, safe="/%:@-._~!$&'()*+,;=")
    query = urllib.parse.quote(parts.query, safe="=&%+;,:/?@-._~!$'()*")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, parts.fragment))


def request_bytes(url: str, *, timeout: float = 60.0) -> tuple[bytes, str]:
    request = urllib.request.Request(
        sanitize_url(url),
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.5",
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read(), response.geturl()


def sha256_hex(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def is_numeric(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def normalize_text(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    text = re.sub(r"\s+", " ", value).strip()
    if not text or len(text) > 240:
        return None
    return text


def label_categories(text: str) -> list[str]:
    return [
        group
        for group, patterns in COMPILED_PATTERNS.items()
        if any(pattern.search(text) for pattern in patterns)
    ]


def numeric_run(
    sheet: Any,
    row: int,
    column: int,
    *,
    direction: str,
    maximum_steps: int = 20000,
) -> dict[str, Any] | None:
    values: list[float] = []
    coordinates: list[str] = []
    blanks_after_start = 0
    for step in range(1, maximum_steps + 1):
        target_row = row + step if direction == "vertical" else row
        target_column = column if direction == "vertical" else column + step
        if target_row > sheet.max_row or target_column > sheet.max_column:
            break
        cell = sheet.cell(target_row, target_column)
        value = cell.value
        if is_numeric(value):
            values.append(float(value))
            coordinates.append(cell.coordinate)
            blanks_after_start = 0
            continue
        if values and value in (None, ""):
            blanks_after_start += 1
            if blanks_after_start <= 1:
                continue
        break
    if len(values) < 2:
        return None
    return {
        "direction": direction,
        "count": len(values),
        "distinct_count": len(set(values)),
        "minimum": min(values),
        "maximum": max(values),
        "first_coordinate": coordinates[0],
        "last_coordinate": coordinates[-1],
    }


def connected_numeric_blocks(
    numeric_values: dict[tuple[int, int], float],
    *,
    minimum_block_size: int = 3,
    maximum_blocks: int = 80,
) -> tuple[list[dict[str, Any]], int]:
    remaining = set(numeric_values)
    blocks: list[dict[str, Any]] = []
    isolated_or_small = 0
    while remaining:
        start = remaining.pop()
        queue: deque[tuple[int, int]] = deque([start])
        component = [start]
        while queue:
            row, column = queue.popleft()
            for neighbour in (
                (row - 1, column),
                (row + 1, column),
                (row, column - 1),
                (row, column + 1),
            ):
                if neighbour in remaining:
                    remaining.remove(neighbour)
                    queue.append(neighbour)
                    component.append(neighbour)
        if len(component) < minimum_block_size:
            isolated_or_small += len(component)
            continue
        rows = [item[0] for item in component]
        columns = [item[1] for item in component]
        values = [numeric_values[item] for item in component]
        blocks.append(
            {
                "numeric_cell_count": len(component),
                "minimum_row": min(rows),
                "maximum_row": max(rows),
                "minimum_column": min(columns),
                "maximum_column": max(columns),
                "distinct_row_count": len(set(rows)),
                "distinct_column_count": len(set(columns)),
                "minimum_value": min(values),
                "maximum_value": max(values),
            }
        )
    blocks.sort(key=lambda item: item["numeric_cell_count"], reverse=True)
    return blocks[:maximum_blocks], isolated_or_small


def infer_layout(
    numeric_blocks: list[dict[str, Any]],
    labeled_runs: list[dict[str, Any]],
    numeric_cell_count: int,
) -> str:
    temperature_runs = [item for item in labeled_runs if "temperature" in item["categories"]]
    density_runs = [item for item in labeled_runs if "density" in item["categories"]]
    observable_runs = [
        item
        for item in labeled_runs
        if any(category in item["categories"] for category in ("electrical", "thermal", "derived"))
    ]

    if temperature_runs and density_runs and observable_runs:
        vertical_categories = {
            category
            for item in labeled_runs
            if item["run"]["direction"] == "vertical"
            for category in item["categories"]
        }
        if {"temperature", "density"}.issubset(vertical_categories):
            return "LONG_FORM_TABLE_CANDIDATE"
        if len(temperature_runs) >= 3 and density_runs:
            return "WIDE_TEMPERATURE_SERIES_CANDIDATE"

    if any(
        block["distinct_row_count"] >= 5 and block["distinct_column_count"] >= 5
        for block in numeric_blocks
    ) and temperature_runs and density_runs:
        return "GRID_SURFACE_CANDIDATE"

    if numeric_cell_count < 30:
        return "SUMMARY_OR_DERIVED_TABLE"
    if len(numeric_blocks) >= 2:
        return "MULTIPLE_ONE_DIMENSIONAL_CURVES"
    return "UNRESOLVED_LAYOUT"


def surface_precheck(
    labeled_runs: list[dict[str, Any]],
    text_headers: list[dict[str, Any]],
    numeric_cell_count: int,
    gate: dict[str, Any],
) -> dict[str, Any]:
    temperature_support = max(
        (
            item["run"]["distinct_count"]
            for item in labeled_runs
            if "temperature" in item["categories"]
        ),
        default=0,
    )
    density_support = max(
        (
            item["run"]["distinct_count"]
            for item in labeled_runs
            if "density" in item["categories"]
        ),
        default=0,
    )
    observable_categories = sorted(
        {
            category
            for item in text_headers
            for category in item["categories"]
            if category in {"electrical", "thermal", "derived"}
        }
    )
    units_present = any("units_or_scaling" in item["categories"] for item in text_headers)
    uncertainty_present = any("uncertainty" in item["categories"] for item in text_headers)

    structural_pass = (
        temperature_support >= int(gate["minimum_distinct_temperature_support"])
        and density_support >= int(gate["minimum_distinct_density_support"])
        and numeric_cell_count >= int(gate["minimum_numeric_observations"])
        and bool(observable_categories)
        and units_present
    )
    return {
        "maximum_distinct_temperature_support": temperature_support,
        "maximum_distinct_density_support": density_support,
        "observable_categories": observable_categories,
        "units_or_scaling_note_present": units_present,
        "uncertainty_marker_present": uncertainty_present,
        "structural_surface_gate_pass": structural_pass,
        "fit_allowed": False,
        "fit_blocker": (
            "UNCERTAINTY_OR_REPLICATE_CONTRACT_NOT_YET_FROZEN"
            if structural_pass
            else "STRUCTURAL_SURFACE_GATE_NOT_MET"
        ),
    }


def audit_sheet(sheet: Any, contract: dict[str, Any]) -> dict[str, Any]:
    maximum_headers = int(contract["audit_contract"]["maximum_recorded_text_headers_per_sheet"])
    numeric_values: dict[tuple[int, int], float] = {}
    text_headers: list[dict[str, Any]] = []
    formula_count = 0
    row_counts: Counter[int] = Counter()
    column_counts: Counter[int] = Counter()

    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if cell.data_type == "f":
                formula_count += 1
            if is_numeric(value):
                numeric_values[(cell.row, cell.column)] = float(value)
                row_counts[cell.row] += 1
                column_counts[cell.column] += 1
                continue
            text = normalize_text(value)
            if text is not None and len(text_headers) < maximum_headers:
                text_headers.append(
                    {
                        "coordinate": cell.coordinate,
                        "text": text,
                        "categories": label_categories(text),
                    }
                )

    labeled_runs: list[dict[str, Any]] = []
    for header in text_headers:
        cell = sheet[header["coordinate"]]
        for direction in ("vertical", "horizontal"):
            run = numeric_run(sheet, cell.row, cell.column, direction=direction)
            if run is not None:
                labeled_runs.append(
                    {
                        "coordinate": header["coordinate"],
                        "text": header["text"],
                        "categories": header["categories"],
                        "run": run,
                    }
                )

    numeric_blocks, small_numeric_cells = connected_numeric_blocks(numeric_values)
    layout = infer_layout(numeric_blocks, labeled_runs, len(numeric_values))
    precheck = surface_precheck(
        labeled_runs,
        text_headers,
        len(numeric_values),
        contract["surface_gate"],
    )

    return {
        "title": sheet.title,
        "state": sheet.sheet_state,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "numeric_cell_count": len(numeric_values),
        "formula_cell_count": formula_count,
        "recorded_text_header_count": len(text_headers),
        "text_headers": text_headers,
        "labeled_numeric_runs": labeled_runs,
        "connected_numeric_blocks": numeric_blocks,
        "isolated_or_small_numeric_cell_count": small_numeric_cells,
        "top_numeric_rows": [
            {"row": row, "numeric_cell_count": count}
            for row, count in row_counts.most_common(12)
        ],
        "top_numeric_columns": [
            {"column": column, "numeric_cell_count": count}
            for column, count in column_counts.most_common(12)
        ],
        "layout_class": layout,
        "surface_precheck": precheck,
        "raw_numeric_cells_recorded": False,
        "full_numeric_arrays_recorded": False,
    }


def audit_workbook(
    workbook_meta: dict[str, Any],
    target_sheets: Iterable[str],
    contract: dict[str, Any],
) -> dict[str, Any]:
    from openpyxl import load_workbook

    filename = str(workbook_meta["filename"])
    payload, resolved_url = request_bytes(str(contract["download_base"]) + filename)
    observed_hash = sha256_hex(payload)
    errors: list[str] = []
    if observed_hash != workbook_meta["sha256"]:
        errors.append(
            f"sha256 changed: expected {workbook_meta['sha256']} observed {observed_hash}"
        )
    if len(payload) != int(workbook_meta["byte_size"]):
        errors.append(
            f"byte size changed: expected {workbook_meta['byte_size']} observed {len(payload)}"
        )

    workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    sheets: list[dict[str, Any]] = []
    for title in target_sheets:
        if title not in workbook.sheetnames:
            errors.append(f"target sheet missing: {title}")
            continue
        sheets.append(audit_sheet(workbook[title], contract))
    workbook.close()

    return {
        "filename": filename,
        "resolved_url": resolved_url,
        "byte_size": len(payload),
        "sha256": observed_hash,
        "pin_pass": not errors,
        "errors": errors,
        "sheets": sheets,
        "raw_workbook_stored": False,
    }


def run_contract(contract: dict[str, Any], source_certificate: dict[str, Any]) -> dict[str, Any]:
    metadata = {
        item["filename"]: item
        for item in source_certificate.get("workbooks", [])
        if isinstance(item, dict) and isinstance(item.get("filename"), str)
    }
    workbook_results: list[dict[str, Any]] = []
    for filename, target_sheets in contract["target_sheets"].items():
        if filename not in metadata:
            workbook_results.append(
                {
                    "filename": filename,
                    "pin_pass": False,
                    "errors": ["workbook absent from G05 source certificate"],
                    "sheets": [],
                    "raw_workbook_stored": False,
                }
            )
            continue
        try:
            workbook_results.append(
                audit_workbook(metadata[filename], target_sheets, contract)
            )
        except Exception as exc:
            workbook_results.append(
                {
                    "filename": filename,
                    "pin_pass": False,
                    "errors": [f"{type(exc).__name__}: {exc}"],
                    "sheets": [],
                    "raw_workbook_stored": False,
                }
            )

    all_sheets = [sheet for book in workbook_results for sheet in book.get("sheets", [])]
    structural_candidates = [
        {
            "filename": book["filename"],
            "sheet": sheet["title"],
            "layout_class": sheet["layout_class"],
            "surface_precheck": sheet["surface_precheck"],
        }
        for book in workbook_results
        for sheet in book.get("sheets", [])
        if sheet["surface_precheck"]["structural_surface_gate_pass"]
    ]
    pin_failures = [book["filename"] for book in workbook_results if not book.get("pin_pass")]
    status = "PASS_NUMERICAL_TOPOLOGY" if not pin_failures else "FAIL_NUMERICAL_TOPOLOGY"

    return {
        "campaign": contract["campaign"],
        "doi": contract["doi"],
        "status": status,
        "workbook_count": len(workbook_results),
        "sheet_count": len(all_sheets),
        "structural_surface_candidate_count": len(structural_candidates),
        "structural_surface_candidates": structural_candidates,
        "pin_failures": pin_failures,
        "workbooks": workbook_results,
        "raw_workbooks_stored": False,
        "raw_numeric_cells_recorded": False,
        "experimental_plucker_significance_computed": False,
        "claim_boundary": contract["claim_boundary"],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "contract",
        nargs="?",
        type=Path,
        default=Path("protocols/G07_SOURCE_DATA_TOPOLOGY.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/G07_NUMERICAL_TOPOLOGY.json"),
    )
    args = parser.parse_args()

    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    source_certificate = json.loads(
        Path(contract["source_certificate"]).read_text(encoding="utf-8")
    )
    result = run_contract(contract, source_certificate)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0 if result["status"] == "PASS_NUMERICAL_TOPOLOGY" else 1


if __name__ == "__main__":
    raise SystemExit(main())
