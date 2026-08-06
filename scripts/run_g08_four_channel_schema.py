from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

from openpyxl.utils import column_index_from_string


USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "Chrome/151.0 Safari/537.36 Thermodynamics-Reproducibility-G08/1.0"
)


def sanitize_url(url: str) -> str:
    parts = urllib.parse.urlsplit(url.strip())
    path = urllib.parse.quote(parts.path, safe="/%:@-._~!$&'()*+,;=")
    query = urllib.parse.quote(parts.query, safe="=&%+;,:/?@-._~!$'()*")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, parts.fragment))


def request_bytes(url: str, *, timeout: float = 60.0) -> tuple[bytes, str]:
    request = urllib.request.Request(
        sanitize_url(url),
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.5",
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read(), response.geturl()


def sha256_hex(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def is_numeric(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def parse_temperature(value: object) -> float | None:
    if is_numeric(value):
        return float(value)
    if isinstance(value, str):
        match = re.search(r"(-?\d+(?:\.\d+)?)\s*K\b", value, flags=re.I)
        if match:
            return float(match.group(1))
    return None


def choose_temperature(values: list[object]) -> float | None:
    parsed = [temperature for value in values if (temperature := parse_temperature(value)) is not None]
    if not parsed:
        return None
    first = parsed[0]
    if any(not math.isclose(first, value, rel_tol=0.0, abs_tol=1e-9) for value in parsed[1:]):
        raise ValueError(f"inconsistent temperature metadata: {parsed}")
    return first


def paired_series_audit(sheet: Any, channel: dict[str, Any]) -> dict[str, Any]:
    series: list[dict[str, Any]] = []
    errors: list[str] = []
    scale = float(channel["density_scale_to_chart_unit"])
    metadata_row = int(channel["temperature_metadata_row"])
    unit_row = int(channel["unit_row"])
    data_start = int(channel["data_start_row"])
    minimum_points = 5

    for pair_index, pair in enumerate(channel["pair_columns"]):
        x_letter, y_letter = pair
        x_column = column_index_from_string(x_letter)
        y_column = column_index_from_string(y_letter)
        x_header = sheet.cell(1, x_column).value
        y_header = sheet.cell(1, y_column).value

        if str(x_header).strip() != str(channel["x_header"]):
            errors.append(
                f"pair {pair_index}: x header changed at {x_letter}1: {x_header!r}"
            )
        if str(y_header).strip() != str(channel["y_header"]):
            errors.append(
                f"pair {pair_index}: y header changed at {y_letter}1: {y_header!r}"
            )

        metadata_column = channel.get("temperature_metadata_column")
        if metadata_column == "x":
            metadata_values = [sheet.cell(metadata_row, x_column).value]
        elif metadata_column == "y":
            metadata_values = [sheet.cell(metadata_row, y_column).value]
        else:
            metadata_values = [
                sheet.cell(metadata_row, x_column).value,
                sheet.cell(metadata_row, y_column).value,
            ]
        try:
            temperature = choose_temperature(metadata_values)
        except ValueError as exc:
            errors.append(f"pair {pair_index}: {exc}")
            temperature = None
        if temperature is None:
            errors.append(f"pair {pair_index}: temperature metadata unresolved")

        x_values: list[float] = []
        y_values: list[float] = []
        rows: list[int] = []
        missing_pair_rows = 0
        for row in range(data_start, sheet.max_row + 1):
            x_value = sheet.cell(row, x_column).value
            y_value = sheet.cell(row, y_column).value
            x_numeric = is_numeric(x_value)
            y_numeric = is_numeric(y_value)
            if x_numeric and y_numeric:
                x_values.append(float(x_value) * scale)
                y_values.append(float(y_value))
                rows.append(row)
            elif x_numeric or y_numeric:
                missing_pair_rows += 1

        if len(x_values) < minimum_points:
            errors.append(
                f"pair {pair_index}: only {len(x_values)} complete points, expected >= {minimum_points}"
            )
        if not x_values:
            continue

        series.append(
            {
                "pair_index": pair_index,
                "x_column": x_letter,
                "y_column": y_letter,
                "temperature_K": temperature,
                "x_header": x_header,
                "y_header": y_header,
                "x_unit": sheet.cell(unit_row, x_column).value,
                "y_unit": sheet.cell(unit_row, y_column).value,
                "data_start_row": min(rows),
                "data_end_row": max(rows),
                "complete_point_count": len(x_values),
                "missing_pair_row_count": missing_pair_rows,
                "distinct_density_count": len(set(x_values)),
                "density_min_1e12_cm_minus_2": min(x_values),
                "density_max_1e12_cm_minus_2": max(x_values),
                "observable_min": min(y_values),
                "observable_max": max(y_values),
            }
        )

    temperatures = sorted(
        {float(item["temperature_K"]) for item in series if item["temperature_K"] is not None}
    )
    density_common_min = max(
        (item["density_min_1e12_cm_minus_2"] for item in series),
        default=None,
    )
    density_common_max = min(
        (item["density_max_1e12_cm_minus_2"] for item in series),
        default=None,
    )
    density_common_nonempty = (
        density_common_min is not None
        and density_common_max is not None
        and density_common_min <= density_common_max
    )

    return {
        "channel_id": channel["channel_id"],
        "observable": channel["observable"],
        "schema_type": channel["schema_type"],
        "measurement_class": channel["measurement_class"],
        "reported_uncertainty_in_sheet": channel["reported_uncertainty_in_sheet"],
        "series_count": len(series),
        "series": series,
        "temperature_support_K": temperatures,
        "temperature_range_K": [min(temperatures), max(temperatures)] if temperatures else None,
        "channel_common_density_range_1e12_cm_minus_2": (
            [density_common_min, density_common_max] if density_common_nonempty else None
        ),
        "schema_pass": not errors,
        "errors": errors,
        "raw_values_stored": False,
    }


def grid_audit(sheet: Any, channel: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    scale = float(channel["density_scale_to_chart_unit"])
    temperature_column = column_index_from_string(channel["temperature_column"])
    temperature_rows = range(
        int(channel["temperature_start_row"]),
        int(channel["temperature_end_row"]) + 1,
    )
    density_row = int(channel["density_metadata_row"])
    density_columns = range(
        column_index_from_string(channel["density_start_column"]),
        column_index_from_string(channel["density_end_column"]) + 1,
    )

    temperatures: list[float] = []
    for row in temperature_rows:
        value = sheet.cell(row, temperature_column).value
        if not is_numeric(value):
            errors.append(f"non-numeric temperature at row {row}: {value!r}")
        else:
            temperatures.append(float(value))

    densities: list[float] = []
    for column in density_columns:
        value = sheet.cell(density_row, column).value
        if not is_numeric(value):
            errors.append(f"non-numeric density at column {column}: {value!r}")
        else:
            densities.append(float(value) * scale)

    observable_values: list[float] = []
    missing_cells = 0
    for row in range(
        int(channel["observable_start_row"]),
        int(channel["observable_end_row"]) + 1,
    ):
        for column in range(
            column_index_from_string(channel["observable_start_column"]),
            column_index_from_string(channel["observable_end_column"]) + 1,
        ):
            value = sheet.cell(row, column).value
            if is_numeric(value):
                observable_values.append(float(value))
            else:
                missing_cells += 1

    expected_cells = len(temperatures) * len(densities)
    if len(observable_values) + missing_cells != expected_cells:
        errors.append("observable grid dimensions are inconsistent")

    return {
        "channel_id": channel["channel_id"],
        "observable": channel["observable"],
        "schema_type": channel["schema_type"],
        "measurement_class": channel["measurement_class"],
        "reported_uncertainty_in_sheet": channel["reported_uncertainty_in_sheet"],
        "temperature_support_K": sorted(set(temperatures)),
        "temperature_range_K": [min(temperatures), max(temperatures)] if temperatures else None,
        "density_support_1e12_cm_minus_2": sorted(set(densities)),
        "channel_common_density_range_1e12_cm_minus_2": (
            [min(densities), max(densities)] if densities else None
        ),
        "grid_shape": [len(temperatures), len(densities)],
        "complete_observation_count": len(observable_values),
        "missing_observation_count": missing_cells,
        "observable_min": min(observable_values) if observable_values else None,
        "observable_max": max(observable_values) if observable_values else None,
        "schema_pass": not errors,
        "errors": errors,
        "raw_values_stored": False,
    }


def audit_channel(
    channel: dict[str, Any],
    workbook_payload: bytes,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(workbook_payload), read_only=False, data_only=False)
    try:
        if channel["sheet"] not in workbook.sheetnames:
            return {
                "channel_id": channel["channel_id"],
                "schema_pass": False,
                "errors": [f"missing sheet: {channel['sheet']}"],
                "raw_values_stored": False,
            }
        sheet = workbook[channel["sheet"]]
        if channel["schema_type"] == "PAIRED_WIDE_SERIES":
            return paired_series_audit(sheet, channel)
        if channel["schema_type"] == "TEMPERATURE_BY_DENSITY_GRID":
            return grid_audit(sheet, channel)
        return {
            "channel_id": channel["channel_id"],
            "schema_pass": False,
            "errors": [f"unsupported schema type: {channel['schema_type']}"],
            "raw_values_stored": False,
        }
    finally:
        workbook.close()


def interval_intersection(intervals: list[list[float] | None]) -> list[float] | None:
    valid = [interval for interval in intervals if interval is not None]
    if len(valid) != len(intervals) or not valid:
        return None
    lower = max(float(interval[0]) for interval in valid)
    upper = min(float(interval[1]) for interval in valid)
    return [lower, upper] if lower <= upper else None


def run_contract(contract: dict[str, Any], source_certificate: dict[str, Any]) -> dict[str, Any]:
    source_meta = {
        item["filename"]: item
        for item in source_certificate.get("workbooks", [])
        if isinstance(item, dict) and isinstance(item.get("filename"), str)
    }
    payload_cache: dict[str, bytes] = {}
    workbook_pins: list[dict[str, Any]] = []
    pin_errors: list[str] = []

    for filename in sorted({channel["workbook"] for channel in contract["channels"]}):
        if filename not in source_meta:
            pin_errors.append(f"{filename}: absent from G05 certificate")
            continue
        try:
            payload, resolved_url = request_bytes(contract["download_base"] + filename)
            digest = sha256_hex(payload)
            expected = source_meta[filename]
            errors: list[str] = []
            if digest != expected["sha256"]:
                errors.append(
                    f"sha256 changed: expected {expected['sha256']} observed {digest}"
                )
            if len(payload) != int(expected["byte_size"]):
                errors.append(
                    f"byte size changed: expected {expected['byte_size']} observed {len(payload)}"
                )
            if not errors:
                payload_cache[filename] = payload
            else:
                pin_errors.extend(f"{filename}: {error}" for error in errors)
            workbook_pins.append(
                {
                    "filename": filename,
                    "resolved_url": resolved_url,
                    "byte_size": len(payload),
                    "sha256": digest,
                    "pin_pass": not errors,
                    "errors": errors,
                }
            )
        except Exception as exc:
            pin_errors.append(f"{filename}: {type(exc).__name__}: {exc}")

    channel_results: list[dict[str, Any]] = []
    for channel in contract["channels"]:
        payload = payload_cache.get(channel["workbook"])
        if payload is None:
            channel_results.append(
                {
                    "channel_id": channel["channel_id"],
                    "schema_pass": False,
                    "errors": ["workbook pin failed or payload unavailable"],
                    "raw_values_stored": False,
                }
            )
            continue
        channel_results.append(audit_channel(channel, payload))

    schema_errors = [
        f"{result.get('channel_id')}: {error}"
        for result in channel_results
        for error in result.get("errors", [])
    ]
    temperature_sets = [
        {round(float(value), 9) for value in result.get("temperature_support_K", [])}
        for result in channel_results
    ]
    exact_temperature_intersection = sorted(set.intersection(*temperature_sets)) if temperature_sets and all(temperature_sets) else []
    common_temperature_range = interval_intersection(
        [result.get("temperature_range_K") for result in channel_results]
    )
    common_density_range = interval_intersection(
        [
            result.get("channel_common_density_range_1e12_cm_minus_2")
            for result in channel_results
        ]
    )
    strict_common_domain = common_temperature_range is not None and common_density_range is not None
    uncertainty_ready = all(
        bool(result.get("reported_uncertainty_in_sheet")) for result in channel_results
    )
    independence_ready = bool(
        contract["independence_contract"]["independent_six_bracket_falsification"]
    )

    if pin_errors or schema_errors:
        status = "FAIL_SCHEMA_AUDIT"
    elif not strict_common_domain:
        status = "INCONCLUSIVE_COMMON_DOMAIN"
    elif not uncertainty_ready:
        status = "INCONCLUSIVE_UNCERTAINTY_MODEL"
    elif not independence_ready:
        status = "INCONCLUSIVE_SOURCE_INDEPENDENCE"
    else:
        status = "PASS_SCHEMA_AUDIT"

    return {
        "campaign": contract["campaign"],
        "doi": contract["doi"],
        "status": status,
        "schema_audit_pass": not pin_errors and not schema_errors,
        "workbook_pins": workbook_pins,
        "pin_errors": pin_errors,
        "schema_errors": schema_errors,
        "channels": channel_results,
        "coverage": {
            "exact_temperature_support_intersection_K": exact_temperature_intersection,
            "continuous_temperature_range_intersection_K": common_temperature_range,
            "density_range_intersection_1e12_cm_minus_2": common_density_range,
            "strict_common_domain_nonempty": strict_common_domain,
        },
        "uncertainty_contract_ready": uncertainty_ready,
        "independent_six_bracket_falsification_ready": independence_ready,
        "fit_allowed": False,
        "fit_blockers": [
            blocker
            for blocker, active in (
                ("SCHEMA_OR_PIN_FAILURE", bool(pin_errors or schema_errors)),
                ("EMPTY_COMMON_DOMAIN", not strict_common_domain),
                ("UNCERTAINTY_MODEL_NOT_FROZEN", not uncertainty_ready),
                ("SOURCE_INDEPENDENCE_NOT_MET", not independence_ready),
            )
            if active
        ],
        "excluded_model_columns": {
            channel["channel_id"]: channel.get("excluded_model_columns")
            for channel in contract["channels"]
            if channel.get("excluded_model_columns")
        },
        "raw_workbooks_stored": False,
        "raw_values_stored": False,
        "experimental_plucker_significance_computed": False,
        "claim_boundary": contract["claim_boundary"],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "contract",
        nargs="?",
        type=Path,
        default=Path("protocols/G08_FOUR_CHANNEL_SCHEMA.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/G08_FOUR_CHANNEL_SCHEMA_AUDIT.json"),
    )
    args = parser.parse_args()

    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    source_certificate = json.loads(
        Path(contract["source_certificate"]).read_text(encoding="utf-8")
    )
    result = run_contract(contract, source_certificate)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, sort_keys=True))
    return 1 if result["status"] == "FAIL_SCHEMA_AUDIT" else 0


if __name__ == "__main__":
    raise SystemExit(main())
