from __future__ import annotations

import argparse
import hashlib
import html
import io
import json
import re
import urllib.parse
import urllib.request
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "Chrome/151.0 Safari/537.36 Thermodynamics-Reproducibility-G05/1.0"
)
ALLOWED_HOST_SUFFIXES = (
    "nature.com",
    "springernature.com",
    "springer.com",
)


class LinkCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        for key, value in attrs:
            if key.lower() in {"href", "src", "data-url", "data-download-url"} and value:
                self.links.append(html.unescape(value))


def sanitize_url(url: str) -> str:
    parts = urllib.parse.urlsplit(html.unescape(url.strip()))
    path = urllib.parse.quote(parts.path, safe="/%:@-._~!$&'()*+,;=")
    query = urllib.parse.quote(parts.query, safe="=&%+;,:/?@-._~!$'()*")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, parts.fragment))


def allowed_host(url: str) -> bool:
    try:
        host = urllib.parse.urlsplit(url).hostname
    except ValueError:
        return False
    return bool(host and any(host == suffix or host.endswith("." + suffix) for suffix in ALLOWED_HOST_SUFFIXES))


def request_bytes(url: str, *, timeout: float = 35.0) -> tuple[bytes, str, str]:
    clean = sanitize_url(url)
    request = urllib.request.Request(
        clean,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,"
                "application/vnd.ms-excel,application/zip,text/html;q=0.8,*/*;q=0.5"
            ),
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read(), response.geturl(), response.headers.get_content_type()


def discover_article_links(article_url: str, payload: bytes) -> list[str]:
    text = payload.decode("utf-8", errors="replace")
    parser = LinkCollector()
    parser.feed(text)
    regex_links = re.findall(r"https?://[^\"'<>\s]+", html.unescape(text))
    links: list[str] = []
    for raw in parser.links + regex_links:
        try:
            resolved = sanitize_url(urllib.parse.urljoin(article_url, raw))
        except (TypeError, ValueError, UnicodeError):
            continue
        lower = urllib.parse.unquote(resolved).casefold()
        if allowed_host(resolved) and any(
            token in lower
            for token in ("source-data", ".xlsx", ".xls", ".zip", "moesm")
        ):
            links.append(resolved)
    return list(dict.fromkeys(links))


def candidate_urls(contract: dict[str, Any], discovered: list[str]) -> list[str]:
    urls = list(discovered)
    stem = str(contract["media_stem"])
    for base in contract["candidate_bases"]:
        for index in contract["candidate_indices"]:
            for extension in contract["candidate_extensions"]:
                urls.append(f"{base}{stem}_MOESM{index}_ESM.{extension}")
    return list(dict.fromkeys(sanitize_url(url) for url in urls if allowed_host(url)))


def sha256_hex(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def is_xlsx(payload: bytes) -> bool:
    if not payload.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = set(archive.namelist())
        return "[Content_Types].xml" in names and "xl/workbook.xml" in names
    except zipfile.BadZipFile:
        return False


def extract_xlsx_members(payload: bytes) -> list[tuple[str, bytes]]:
    if is_xlsx(payload):
        return [("workbook.xlsx", payload)]
    if not payload.startswith(b"PK"):
        return []
    results: list[tuple[str, bytes]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            for name in archive.namelist():
                if Path(name).suffix.lower() == ".xlsx":
                    member = archive.read(name)
                    if is_xlsx(member):
                        results.append((Path(name).name, member))
    except zipfile.BadZipFile:
        return []
    return results


def workbook_audit(filename: str, payload: bytes, resolved_url: str) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        numeric_cells = 0
        text_cells = 0
        formula_cells = 0
        blank_cells = 0
        error_cells = 0
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    blank_cells += 1
                elif cell.data_type == "f":
                    formula_cells += 1
                elif cell.data_type == "e":
                    error_cells += 1
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric_cells += 1
                else:
                    text_cells += 1
        sheets.append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "numeric_cell_count": numeric_cells,
                "text_cell_count": text_cells,
                "formula_cell_count": formula_cells,
                "blank_cell_count": blank_cells,
                "error_cell_count": error_cells,
                "merged_range_count": len(sheet.merged_cells.ranges),
            }
        )
    workbook.close()
    return {
        "filename": filename,
        "resolved_url": resolved_url,
        "byte_size": len(payload),
        "sha256": sha256_hex(payload),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "raw_workbook_stored": False,
        "raw_workbook_uploaded_as_artifact": False,
    }


def run_contract(contract: dict[str, Any]) -> dict[str, Any]:
    attempts: list[dict[str, str]] = []
    discovered: list[str] = []
    try:
        page, final_url, content_type = request_bytes(str(contract["article_url"]))
        attempts.append(
            {
                "url": str(contract["article_url"]),
                "status": "ARTICLE_FETCHED",
                "content_type": content_type,
            }
        )
        discovered = discover_article_links(final_url, page)
    except Exception as exc:
        attempts.append(
            {
                "url": str(contract["article_url"]),
                "status": "ARTICLE_FETCH_FAILED",
                "error": f"{type(exc).__name__}: {exc}",
            }
        )

    workbooks: list[dict[str, Any]] = []
    seen_hashes: set[str] = set()
    for url in candidate_urls(contract, discovered):
        try:
            payload, final_url, content_type = request_bytes(url)
            members = extract_xlsx_members(payload)
            if not members:
                attempts.append(
                    {
                        "url": url,
                        "status": "NOT_A_WORKBOOK",
                        "content_type": content_type,
                    }
                )
                continue
            for member_name, workbook_payload in members:
                digest = sha256_hex(workbook_payload)
                if digest in seen_hashes:
                    continue
                seen_hashes.add(digest)
                workbooks.append(workbook_audit(member_name, workbook_payload, final_url))
            attempts.append({"url": url, "status": "WORKBOOK_FOUND"})
        except Exception as exc:
            attempts.append(
                {
                    "url": url,
                    "status": "FAILED",
                    "error": f"{type(exc).__name__}: {exc}",
                }
            )

    status = (
        "PASS_SOURCE_DATA_DISCOVERY"
        if workbooks
        else "INCONCLUSIVE_SOURCE_DATA_DISCOVERY"
    )
    return {
        "campaign": contract["campaign"],
        "doi": contract["doi"],
        "status": status,
        "discovered_article_link_count": len(discovered),
        "workbook_count": len(workbooks),
        "workbooks": workbooks,
        "attempts": attempts,
        "raw_workbooks_stored": False,
        "raw_workbooks_uploaded_as_artifacts": False,
        "experimental_plucker_significance_computed": False,
        "claim_boundary": contract["claim_boundary"],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "contract",
        nargs="?",
        type=Path,
        default=Path("protocols/G05_MAJUMDAR_SOURCE_DATA.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/G05_SOURCE_DATA_DISCOVERY.json"),
    )
    args = parser.parse_args()

    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    result = run_contract(contract)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
