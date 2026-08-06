from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any


USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "Chrome/151.0 Safari/537.36 Thermodynamics-Reproducibility-G06/1.0"
)


def sanitize_url(url: str) -> str:
    parts = urllib.parse.urlsplit(url.strip())
    path = urllib.parse.quote(parts.path, safe="/%:@-._~!$&'()*+,;=")
    query = urllib.parse.quote(parts.query, safe="=&%+;,:/?@-._~!$'()*")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, parts.fragment))


def request_bytes(url: str, *, timeout: float = 45.0) -> tuple[bytes, str]:
    request = urllib.request.Request(
        sanitize_url(url),
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.5",
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read(), response.geturl()


def sha256_hex(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def normalized_text(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    text = re.sub(r"\s+", " ", value).strip()
    if not text or len(text) > 240:
        return None
    return text


def matched_labels(
    labels: list[str],
    keyword_groups: dict[str, list[str]],
) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for group, keywords in keyword_groups.items():
        matches: list[str] = []
        for label in labels:
            lower = f" {label.casefold()} "
            if any(keyword.casefold() in lower for keyword in keywords):
                matches.append(label)
        result[group] = list(dict.fromkeys(matches))[:80]
    return result


def infer_observables(matches: dict[str, list[str]]) -> list[str]:
    observables: list[str] = []
    if matches.get("electrical_transport"):
        observables.append("electrical_transport")
    if matches.get("thermal_transport"):
        observables.append("thermal_transport")
    if matches.get("derived_hydrodynamics"):
        observables.append("derived_hydrodynamics")
    return observables


def classify_sheet(matches: dict[str, list[str]]) -> tuple[str, list[str]]:
    has_temperature = bool(matches.get("temperature_control"))
    has_density = bool(matches.get("density_control"))
    observables = infer_observables(matches)

    if has_temperature and has_density and observables:
        return "TWO_CONTROL_CHANNEL_CANDIDATE", observables
    if has_temperature and has_density:
        return "TWO_CONTROL_OBSERVABLE_UNRESOLVED", observables
    if observables == ["derived_hydrodynamics"]:
        return "DERIVED_OR_AUXILIARY_CHANNEL", observables
    return "ONE_CONTROL_OR_UNRESOLVED", observables


def audit_workbook(
    workbook_meta: dict[str, Any],
    download_base: str,
    keyword_groups: dict[str, list[str]],
) -> dict[str, Any]:
    from openpyxl import load_workbook

    filename = str(workbook_meta["filename"])
    payload, resolved_url = request_bytes(download_base + filename)
    digest = sha256_hex(payload)
    errors: list[str] = []
    if digest != workbook_meta["sha256"]:
        errors.append(
            f"sha256 changed: expected {workbook_meta['sha256']} observed {digest}"
        )
    if len(payload) != int(workbook_meta["byte_size"]):
        errors.append(
            f"byte size changed: expected {workbook_meta['byte_size']} observed {len(payload)}"
        )

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        labels: list[str] = []
        for row in sheet.iter_rows():
            for cell in row:
                label = normalized_text(cell.value)
                if label is not None:
                    labels.append(label)
        unique_labels = list(dict.fromkeys(labels))
        matches = matched_labels(unique_labels, keyword_groups)
        status, observable_classes = classify_sheet(matches)
        sheets.append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "unique_text_label_count": len(unique_labels),
                "matched_labels": matches,
                "temperature_marker_present": bool(matches["temperature_control"]),
                "density_marker_present": bool(matches["density_control"]),
                "observable_classes": observable_classes,
                "status": status,
                "numeric_cell_values_stored": false_value(),
            }
        )
    workbook.close()

    return {
        "filename": filename,
        "resolved_url": resolved_url,
        "byte_size": len(payload),
        "sha256": digest,
        "hash_pin_pass": not errors,
        "errors": errors,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "raw_workbook_stored": False,
        "numeric_cell_values_stored": False,
    }


def false_value() -> bool:
    return False


def run_contract(
    contract: dict[str, Any],
    source_certificate: dict[str, Any],
) -> dict[str, Any]:
    workbook_results: list[dict[str, Any]] = []
    for workbook_meta in source_certificate.get("workbooks", []):
        try:
            workbook_results.append(
                audit_workbook(
                    workbook_meta,
                    str(contract["download_base"]),
                    contract["keyword_groups"],
                )
            )
        except Exception as exc:
            workbook_results.append(
                {
                    "filename": workbook_meta.get("filename"),
                    "hash_pin_pass": False,
                    "errors": [f"{type(exc).__name__}: {exc}"],
                    "sheets": [],
                    "raw_workbook_stored": False,
                    "numeric_cell_values_stored": False,
                }
            )

    all_sheets = [sheet for book in workbook_results for sheet in book.get("sheets", [])]
    two_control_candidates = [
        {
            "filename": book["filename"],
            "sheet": sheet["title"],
            "observable_classes": sheet["observable_classes"],
        }
        for book in workbook_results
        for sheet in book.get("sheets", [])
        if sheet["status"] == "TWO_CONTROL_CHANNEL_CANDIDATE"
    ]
    pin_failures = [book["filename"] for book in workbook_results if not book.get("hash_pin_pass")]
    status = "PASS_VARIABLE_MAP" if not pin_failures else "FAIL_VARIABLE_MAP"

    return {
        "campaign": contract["campaign"],
        "doi": contract["doi"],
        "status": status,
        "workbook_count": len(workbook_results),
        "sheet_count": len(all_sheets),
        "two_control_candidate_count": len(two_control_candidates),
        "two_control_candidates": two_control_candidates,
        "pin_failures": pin_failures,
        "workbooks": workbook_results,
        "raw_workbooks_stored": False,
        "numeric_cell_values_stored": False,
        "experimental_plucker_significance_computed": False,
        "claim_boundary": contract["claim_boundary"],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "contract",
        nargs="?",
        type=Path,
        default=Path("protocols/G06_SOURCE_DATA_VARIABLE_MAP.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/G06_VARIABLE_MAP.json"),
    )
    args = parser.parse_args()

    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    source_certificate = json.loads(
        Path(contract["source_certificate"]).read_text(encoding="utf-8")
    )
    result = run_contract(contract, source_certificate)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0 if result["status"] == "PASS_VARIABLE_MAP" else 1


if __name__ == "__main__":
    raise SystemExit(main())
