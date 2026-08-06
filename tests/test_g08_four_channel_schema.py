from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook


def load_module():
    script = Path(__file__).parents[1] / "scripts" / "run_g08_four_channel_schema.py"
    spec = importlib.util.spec_from_file_location("g08_schema", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_temperature_parser_accepts_numeric_and_labeled_values() -> None:
    module = load_module()
    assert module.parse_temperature(110) == 110.0
    assert module.parse_temperature("170 K, e doped") == 170.0
    assert module.parse_temperature("not a temperature") is None


def test_paired_series_unit_conversion_and_summary() -> None:
    module = load_module()
    workbook = Workbook()
    sheet = workbook.active
    sheet["B1"] = "n"
    sheet["C1"] = "G"
    sheet["B2"] = 100
    sheet["C2"] = 100
    sheet["B3"] = "m^-2"
    sheet["C3"] = "S"
    for offset, row in enumerate(range(5, 11)):
        sheet.cell(row, 2).value = (offset + 1) * 1e16
        sheet.cell(row, 3).value = offset * 0.1
    channel = {
        "channel_id": "G",
        "observable": "G",
        "schema_type": "PAIRED_WIDE_SERIES",
        "x_header": "n",
        "y_header": "G",
        "temperature_metadata_row": 2,
        "unit_row": 3,
        "data_start_row": 5,
        "density_scale_to_chart_unit": 1e-16,
        "pair_columns": [["B", "C"]],
        "measurement_class": "DIRECT",
        "reported_uncertainty_in_sheet": False,
    }
    result = module.paired_series_audit(sheet, channel)
    workbook.close()
    assert result["schema_pass"] is True
    assert result["temperature_support_K"] == [100.0]
    assert result["series"][0]["density_min_1e12_cm_minus_2"] == 1.0
    assert result["series"][0]["density_max_1e12_cm_minus_2"] == 6.0


def test_grid_schema_extracts_common_axes_without_raw_values() -> None:
    module = load_module()
    workbook = Workbook()
    sheet = workbook.active
    for column, density in zip(range(3, 6), [-0.2, 0.0, 0.2]):
        sheet.cell(3, column).value = density
    for row, temperature in zip(range(5, 8), [100, 150, 200]):
        sheet.cell(row, 2).value = temperature
        for column in range(3, 6):
            sheet.cell(row, column).value = row + column
    channel = {
        "channel_id": "SIGMA",
        "observable": "Sigma",
        "schema_type": "TEMPERATURE_BY_DENSITY_GRID",
        "temperature_column": "B",
        "temperature_start_row": 5,
        "temperature_end_row": 7,
        "density_metadata_row": 3,
        "density_start_column": "C",
        "density_end_column": "E",
        "observable_start_row": 5,
        "observable_end_row": 7,
        "observable_start_column": "C",
        "observable_end_column": "E",
        "density_scale_to_chart_unit": 1.0,
        "measurement_class": "FIT",
        "reported_uncertainty_in_sheet": False,
    }
    result = module.grid_audit(sheet, channel)
    workbook.close()
    assert result["schema_pass"] is True
    assert result["grid_shape"] == [3, 3]
    assert result["complete_observation_count"] == 9
    assert result["raw_values_stored"] is False


def test_interval_intersection_detects_empty_temperature_domain() -> None:
    module = load_module()
    assert module.interval_intersection([[40, 100], [110, 260]]) is None
    assert module.interval_intersection([[40, 170], [110, 260]]) == [110.0, 170.0]
